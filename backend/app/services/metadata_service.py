import os
from datetime import datetime
from typing import Dict, Any, Optional, Tuple
from PIL import Image
from PIL.ExifTags import TAGS, GPSTAGS
from pypdf import PdfReader
import docx
import pptx
import openpyxl

class MetadataService:
    @staticmethod
    def extract_metadata(file_path: str, file_type: str) -> Dict[str, Any]:
        """
        Main entry point to route metadata extraction based on file type.
        """
        result = {
            "creator": None,
            "software_used": None,
            "created_datetime": None,
            "modified_datetime": None,
            "gps_latitude": None,
            "gps_longitude": None,
            "raw_metadata": {}
        }
        
        try:
            if file_type == "image":
                return MetadataService._extract_image_metadata(file_path)
            elif file_type == "document":
                ext = os.path.splitext(file_path)[1].lower()
                if ext == ".pdf":
                    return MetadataService._extract_pdf_metadata(file_path)
                elif ext == ".docx":
                    return MetadataService._extract_docx_metadata(file_path)
                elif ext == ".pptx":
                    return MetadataService._extract_pptx_metadata(file_path)
                elif ext in [".xlsx", ".xls"]:
                    return MetadataService._extract_xlsx_metadata(file_path)
            
            # Default fallback for video/audio/executables/archives: parse system-level metrics
            result["raw_metadata"] = {
                "file_name": os.path.basename(file_path),
                "file_size": os.path.getsize(file_path),
                "last_modified": datetime.fromtimestamp(os.path.getmtime(file_path)).isoformat()
            }
        except Exception as e:
            result["raw_metadata"] = {"error": f"Failed to extract metadata: {str(e)}"}
            
        return result

    @staticmethod
    def _extract_image_metadata(file_path: str) -> Dict[str, Any]:
        res = {
            "creator": None,
            "software_used": None,
            "created_datetime": None,
            "modified_datetime": None,
            "gps_latitude": None,
            "gps_longitude": None,
            "raw_metadata": {}
        }
        
        try:
            with Image.open(file_path) as img:
                res["raw_metadata"]["format"] = img.format
                res["raw_metadata"]["width"] = img.width
                res["raw_metadata"]["height"] = img.height
                res["raw_metadata"]["mode"] = img.mode
                
                exif_data = img._getexif()
                if not exif_data:
                    return res
                
                raw_exif = {}
                gps_info = {}
                for tag, value in exif_data.items():
                    tag_name = TAGS.get(tag, tag)
                    # Convert bytes to string representation if needed
                    if isinstance(value, bytes):
                        try:
                            value = value.decode("utf-8", errors="replace")
                        except Exception:
                            value = str(value)
                    
                    if tag_name == "GPSInfo":
                        for g_tag in value:
                            g_tag_name = GPSTAGS.get(g_tag, g_tag)
                            gps_info[g_tag_name] = str(value[g_tag])
                    else:
                        raw_exif[tag_name] = str(value)
                
                res["raw_metadata"]["exif"] = raw_exif
                res["raw_metadata"]["gps"] = gps_info
                
                # Normalize key fields
                res["software_used"] = raw_exif.get("Software")
                res["creator"] = raw_exif.get("Artist") or raw_exif.get("Copyright")
                
                # Normalize date
                date_str = raw_exif.get("DateTimeOriginal") or raw_exif.get("DateTime")
                if date_str:
                    try:
                        # EXIF format is typically 'YYYY:MM:DD HH:MM:SS'
                        res["created_datetime"] = datetime.strptime(date_str, "%Y:%m:%d %H:%M:%S")
                    except Exception:
                        pass
                
                # Parse GPS Coordinates
                lat, lon = MetadataService._parse_gps_coords(gps_info)
                res["gps_latitude"] = lat
                res["gps_longitude"] = lon
        except Exception as e:
            res["raw_metadata"]["error"] = f"EXIF extraction error: {str(e)}"
            
        return res

    @staticmethod
    def _parse_gps_coords(gps_info: Dict[str, Any]) -> Tuple[Optional[float], Optional[float]]:
        try:
            def to_float(value: Any) -> float:
                if isinstance(value, (int, float)):
                    return float(value)
                if hasattr(value, "numerator") and hasattr(value, "denominator"):
                    denominator = float(value.denominator) or 1.0
                    return float(value.numerator) / denominator
                text = str(value).strip()
                if "/" in text:
                    numerator, denominator = text.split("/", 1)
                    denominator_value = float(denominator) or 1.0
                    return float(numerator) / denominator_value
                return float(text)

            def to_degrees(val: Any) -> float:
                # EXIF GPS values can be tuples, lists, rationals, or strings.
                if isinstance(val, (tuple, list)):
                    parts = [to_float(part) for part in val]
                else:
                    clean = str(val).replace("(", "").replace(")", "").replace("[", "").replace("]", "")
                    parts = [to_float(x.strip()) for x in clean.split(",") if x.strip()]

                if len(parts) >= 3:
                    return parts[0] + parts[1] / 60.0 + parts[2] / 3600.0
                if len(parts) == 1:
                    return parts[0]
                return 0.0

            lat_ref = gps_info.get("GPSLatitudeRef")
            lat_val = gps_info.get("GPSLatitude")
            lon_ref = gps_info.get("GPSLongitudeRef")
            lon_val = gps_info.get("GPSLongitude")
            
            lat = None
            lon = None
            
            if lat_val and lat_ref:
                lat = to_degrees(lat_val)
                if lat_ref != "N":
                    lat = -lat
                    
            if lon_val and lon_ref:
                lon = to_degrees(lon_val)
                if lon_ref != "E":
                    lon = -lon
                    
            return lat, lon
        except Exception:
            return None, None

    @staticmethod
    def _extract_pdf_metadata(file_path: str) -> Dict[str, Any]:
        res = {
            "creator": None,
            "software_used": None,
            "created_datetime": None,
            "modified_datetime": None,
            "gps_latitude": None,
            "gps_longitude": None,
            "raw_metadata": {}
        }
        
        try:
            reader = PdfReader(file_path)
            info = reader.metadata
            if info:
                raw = {}
                for key, val in info.items():
                    raw[key] = str(val)
                res["raw_metadata"] = raw
                
                res["creator"] = info.author or info.creator
                res["software_used"] = info.producer
                
                # PDF timestamps are typically formatted as D:YYYYMMDDHHmmSSOHH'mm'
                def parse_pdf_date(date_str: str) -> Optional[datetime]:
                    if not date_str:
                        return None
                    try:
                        clean_str = date_str.replace("D:", "").split("+")[0].split("-")[0]
                        # Take the first 14 digits YYYYMMDDHHMMSS
                        clean_str = "".join([c for c in clean_str if c.isdigit()])[:14]
                        return datetime.strptime(clean_str, "%Y%m%d%H%M%S")
                    except Exception:
                        return None

                res["created_datetime"] = parse_pdf_date(info.get("/CreationDate"))
                res["modified_datetime"] = parse_pdf_date(info.get("/ModDate"))
                
                res["raw_metadata"]["pages_count"] = len(reader.pages)
        except Exception as e:
            res["raw_metadata"]["error"] = f"PDF extraction error: {str(e)}"
            
        return res

    @staticmethod
    def _extract_docx_metadata(file_path: str) -> Dict[str, Any]:
        res = {
            "creator": None,
            "software_used": None,
            "created_datetime": None,
            "modified_datetime": None,
            "gps_latitude": None,
            "gps_longitude": None,
            "raw_metadata": {}
        }
        
        try:
            doc = docx.Document(file_path)
            props = doc.core_properties
            res["creator"] = props.author or props.last_modified_by
            res["created_datetime"] = props.created
            res["modified_datetime"] = props.modified
            
            res["raw_metadata"] = {
                "title": props.title,
                "subject": props.subject,
                "revision": props.revision,
                "comments": props.comments,
                "version": props.version,
                "last_modified_by": props.last_modified_by
            }
        except Exception as e:
            res["raw_metadata"]["error"] = f"DOCX extraction error: {str(e)}"
            
        return res

    @staticmethod
    def _extract_pptx_metadata(file_path: str) -> Dict[str, Any]:
        res = {
            "creator": None,
            "software_used": None,
            "created_datetime": None,
            "modified_datetime": None,
            "gps_latitude": None,
            "gps_longitude": None,
            "raw_metadata": {}
        }
        
        try:
            prs = pptx.Presentation(file_path)
            props = prs.core_properties
            res["creator"] = props.author or props.last_modified_by
            res["created_datetime"] = props.created
            res["modified_datetime"] = props.modified
            
            res["raw_metadata"] = {
                "title": props.title,
                "subject": props.subject,
                "revision": props.revision,
                "slides_count": len(prs.slides),
                "last_modified_by": props.last_modified_by
            }
        except Exception as e:
            res["raw_metadata"]["error"] = f"PPTX extraction error: {str(e)}"
            
        return res

    @staticmethod
    def _extract_xlsx_metadata(file_path: str) -> Dict[str, Any]:
        res = {
            "creator": None,
            "software_used": None,
            "created_datetime": None,
            "modified_datetime": None,
            "gps_latitude": None,
            "gps_longitude": None,
            "raw_metadata": {}
        }
        
        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            props = wb.properties
            res["creator"] = props.creator or props.lastModifiedBy
            res["created_datetime"] = props.created
            res["modified_datetime"] = props.modified
            
            res["raw_metadata"] = {
                "title": props.title,
                "subject": props.subject,
                "sheets": wb.sheetnames,
                "last_modified_by": props.lastModifiedBy
            }
        except Exception as e:
            res["raw_metadata"]["error"] = f"XLSX extraction error: {str(e)}"
            
        return res
